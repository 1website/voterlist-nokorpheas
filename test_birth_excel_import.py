import io
import sys
import openpyxl
from fastapi.testclient import TestClient

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

from main import app
from app.database import SessionLocal
from app.models import BirthCertificate, Village

client = TestClient(app)

def test_birth_excel_import():
    print("=== Testing Birth Certificates Excel Import ===")

    # 1. Login as Admin
    res_login = client.post("/login", data={"username": "admin", "password": "admin123"}, follow_redirects=False)
    assert res_login.status_code == 303
    cookies = res_login.cookies
    print("[PASS] 1. Admin login authenticated for import test")

    # 2. Download Official Excel Template
    res_tpl = client.get("/birth-certificates/template-excel", cookies=cookies)
    assert res_tpl.status_code == 200
    assert "spreadsheetml" in res_tpl.headers.get("content-type", "")
    
    # Read the returned workbook
    wb = openpyxl.load_workbook(io.BytesIO(res_tpl.content))
    assert "ទម្រង់បញ្ចូលសំបុត្រកំណើត" in wb.sheetnames
    assert "បញ្ជីឈ្មោះភូមិទាំង១០" in wb.sheetnames
    print(f"[PASS] 2. Excel template downloaded and validated ({len(res_tpl.content)} bytes, Sheets: {wb.sheetnames})")

    # 3. Create a test Excel workbook with sample data
    test_wb = openpyxl.Workbook()
    ws = test_wb.active
    ws.title = "ទម្រង់បញ្ចូលសំបុត្រកំណើត"
    headers = [
        "លេខសំបុត្រកំណើត (*)", "សៀវភៅលេខ", "កាលបរិច្ឆេទចុះបញ្ជី", "ឈ្មោះជាភាសាខ្មែរ (*)", 
        "ឈ្មោះជាអក្សរឡាតាំង", "ភេទ (*)", "ថ្ងៃខែឆ្នាំកំណើត (*)", "ទីកន្លែងកំណើត", 
        "ឈ្មោះឪពុក", "ឈ្មោះម្តាយ", "ភូមិ (*)", "អាសយដ្ឋានបច្ចុប្បន្ន", "កំណត់សម្គាល់"
    ]
    ws.append([]) # Row 1
    ws.append([]) # Row 2
    ws.append([]) # Row 3
    ws.append(headers) # Row 4

    test_cert_1 = "ស.ក-2008-TEST01"
    test_cert_2 = "ស.ក-2008-TEST02"

    rows = [
        [test_cert_1, "01/2008", "2026-09-15", "ស៊ឹម តេស្តមួយ", "SIM TESTONE", "ប្រុស", "2008-05-10", "នគរភាស", "ស៊ឹម ឪ", "មាស ម៉ែ", "ភូមិរមៀត", "ក្រុម ០១", "តេស្ត Excel Import"],
        [test_cert_2, "01/2008", "2026-09-15", "មាស តេស្តពីរ", "MEAS TESTTWO", "ស្រី", "2008-08-20", "នគរភាស", "មាស ឪ", "កែវ ម៉ែ", "ភូមិល្បើក", "ក្រុម ០២", "តេស្ត Excel Import 2"]
    ]
    for r in rows:
        ws.append(r)

    excel_file_bytes = io.BytesIO()
    test_wb.save(excel_file_bytes)
    excel_file_bytes.seek(0)

    # 4. Test Preview API
    files = {
        "excel_file": ("test_births.xlsx", excel_file_bytes.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }
    res_prev = client.post("/api/birth-certificates/import-excel/preview", files=files, cookies=cookies)
    p_data = res_prev.json()
    assert p_data["success"] == True
    assert p_data["total_rows"] == 2
    assert p_data["valid_count"] == 2
    assert len(p_data["records"]) == 2
    assert p_data["records"][0]["certificate_no"] == test_cert_1
    assert p_data["records"][0]["name_kh"] == "ស៊ឹម តេស្តមួយ"
    assert p_data["records"][0]["status"] == "valid"
    print(f"[PASS] 3. Preview parsed 2 valid records successfully (Total: {p_data['total_rows']}, Valid: {p_data['valid_count']})")

    # 5. Test Confirm API
    db = SessionLocal()
    try:
        # Clean up if existing from prior test
        db.query(BirthCertificate).filter(BirthCertificate.certificate_no.in_([test_cert_1, test_cert_2])).delete()
        db.commit()
    finally:
        db.close()

    payload = {
        "records": p_data["records"],
        "on_duplicate": "skip"
    }
    res_conf = client.post("/api/birth-certificates/import-excel/confirm", json=payload, cookies=cookies)
    assert res_conf.status_code == 200, f"Confirm failed: {res_conf.text}"
    c_data = res_conf.json()
    assert c_data["success"] == True
    assert c_data["imported_count"] == 2
    print(f"[PASS] 4. Confirm import saved {c_data['imported_count']} records into database")

    # 6. Verify records exist in Database
    db = SessionLocal()
    try:
        b1 = db.query(BirthCertificate).filter(BirthCertificate.certificate_no == test_cert_1).first()
        b2 = db.query(BirthCertificate).filter(BirthCertificate.certificate_no == test_cert_2).first()
        assert b1 is not None
        assert b1.name_kh == "ស៊ឹម តេស្តមួយ"
        assert b1.gender == "ប្រុស"
        assert b1.dob == "2008-05-10"
        assert b1.village.name_kh == "ភូមិរមៀត"
        assert b2 is not None
        assert b2.name_kh == "មាស តេស្តពីរ"
        assert b2.gender == "ស្រី"
        print("[PASS] 5. Database records verified with accurate village, names, gender, and dates")

        # Clean up test records
        db.delete(b1)
        db.delete(b2)
        db.commit()
        print("[PASS] 6. Test records cleaned up successfully")
    finally:
        db.close()

    print("\n ALL BIRTH CERTIFICATES EXCEL IMPORT TESTS PASSED SUCCESSFULLY! ")

if __name__ == "__main__":
    test_birth_excel_import()
